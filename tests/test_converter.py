# -*- coding: utf-8 -*-
"""tests/test_converter.py — Pengujian unit untuk utilitas converter (Excel, CSV, PDF)."""
import io
import os
import sys
import zipfile
import unittest
import openpyxl
import pypdf
from PIL import Image

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
import converter.engine as eng


class ConverterEngineTest(unittest.TestCase):
    def setUp(self):
        # Buat Excel dummy 2 sheet
        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = "Gaji"
        ws1.append(["Nama", "NIK", "Nominal"])
        ws1.append(["Budi", "012345678901234", 15000000])
        ws1.append(["Siti", "098765432109876", 12500000])

        ws2 = wb.create_sheet("Pajak")
        ws2.append(["Kode", "Tarif"])
        ws2.append(["PPH21", 0.05])

        xl_buf = io.BytesIO()
        wb.save(xl_buf)
        self.excel_bytes = xl_buf.getvalue()

        # Buat PDF dummy 3 halaman
        w = pypdf.PdfWriter()
        w.add_blank_page(width=200, height=200)
        w.add_blank_page(width=200, height=200)
        w.add_blank_page(width=200, height=200)
        pdf_buf = io.BytesIO()
        w.write(pdf_buf)
        self.pdf_bytes = pdf_buf.getvalue()

    def test_inspect_excel(self):
        info = eng.inspect_excel(self.excel_bytes)
        self.assertEqual(info["total_sheets"], 2)
        self.assertIn("Gaji", info["sheet_names"])
        self.assertIn("Pajak", info["sheet_names"])
        self.assertGreaterEqual(info["sheets"][0]["estimated_rows"], 3)

    def test_excel_to_csv_single(self):
        csv_bytes, fname, mime, meta = eng.excel_to_csv(
            self.excel_bytes, sheet_name="Gaji", delimiter=";", encoding="utf-8-sig"
        )
        self.assertTrue(fname.endswith(".csv"))
        text = csv_bytes.decode("utf-8-sig")
        lines = text.strip().splitlines()
        self.assertEqual(len(lines), 3)
        self.assertIn("Budi;012345678901234;15000000", lines[1])

    def test_excel_to_csv_all_zip(self):
        zip_bytes, fname, mime, meta = eng.excel_to_csv(
            self.excel_bytes, sheet_name="__all__", delimiter=","
        )
        self.assertTrue(fname.endswith(".zip"))
        self.assertEqual(mime, "application/zip")
        self.assertEqual(meta["total_sheets"], 2)

    def test_inspect_and_convert_csv_to_excel(self):
        csv_data = "NPWP;Nama;Omzet\n012345678901234;PT Maju;500000000\n".encode("utf-8")
        info = eng.inspect_csv(csv_data)
        self.assertEqual(info["detected_delimiter"], ";")

        xl_bytes, fname, mime, meta = eng.csv_to_excel(
            csv_data, delimiter="auto", preserve_text=True
        )
        self.assertTrue(fname.endswith(".xlsx"))
        wb = openpyxl.load_workbook(io.BytesIO(xl_bytes))
        ws = wb.active
        # Pastikan leading zero tetap awet sebagai string
        npwp_val = ws["A2"].value
        self.assertEqual(npwp_val, "012345678901234")
        self.assertIsInstance(npwp_val, str)

    def test_csv_to_csv_reformat(self):
        csv_input = "Nama ; NPWP ; Gaji \n Budi ; 0123456789 ; 15000000 \n\n".encode("utf-8")
        out_b, fname, mime, meta = eng.csv_to_csv(
            csv_input,
            src_delimiter="auto",
            target_delimiter=",",
            target_encoding="utf-8-sig",
            trim_whitespace=True,
            remove_empty_rows=True,
        )
        self.assertTrue(fname.endswith(".csv"))
        text = out_b.decode("utf-8-sig")
        lines = text.strip().splitlines()
        self.assertEqual(len(lines), 2)
        self.assertEqual(lines[0], "Nama,NPWP,Gaji")
        self.assertEqual(lines[1], "Budi,0123456789,15000000")

    def test_pdf_inspect_and_merge(self):
        info = eng.inspect_pdf(self.pdf_bytes)
        self.assertEqual(info["page_count"], 3)
        self.assertFalse(info["is_encrypted"])

        merged_bytes, fname, mime, meta = eng.merge_pdfs([
            ("doc1.pdf", self.pdf_bytes),
            ("doc2.pdf", self.pdf_bytes),
        ])
        self.assertEqual(meta["total_files"], 2)
        self.assertEqual(meta["total_pages"], 6)

        r = pypdf.PdfReader(io.BytesIO(merged_bytes))
        self.assertEqual(len(r.pages), 6)

    def test_pdf_split_all(self):
        zip_bytes, fname, mime, meta = eng.split_pdf(self.pdf_bytes, mode="all")
        self.assertTrue(fname.endswith(".zip"))
        self.assertEqual(meta["files_created"], 3)

    def test_pdf_split_range_merged(self):
        pdf_out, fname, mime, meta = eng.split_pdf(
            self.pdf_bytes, mode="range", range_str="1, 3", merge_range_result=True
        )
        self.assertTrue(fname.endswith(".pdf"))
        self.assertEqual(meta["total_extracted"], 2)
        r = pypdf.PdfReader(io.BytesIO(pdf_out))
        self.assertEqual(len(r.pages), 2)

    def test_pdf_split_range_invalid(self):
        with self.assertRaises(ValueError):
            eng.split_pdf(self.pdf_bytes, mode="range", range_str="1-10")

    def test_pdf_split_step(self):
        zip_bytes, fname, mime, meta = eng.split_pdf(
            self.pdf_bytes, mode="step", step_n=2
        )
        self.assertTrue(fname.endswith(".zip"))
        self.assertEqual(meta["parts_created"], 2)  # hal 1-2, dan hal 3

    def test_excel_html_disguised(self):
        html_bytes = b"<html><body><table><tr><th>ID</th><th>Nama</th></tr><tr><td>1</td><td>Budi</td></tr></table></body></html>"
        info = eng.inspect_excel(html_bytes)
        self.assertEqual(info["engine"], "html")
        self.assertIn("Tabel_1", info["sheet_names"])

        csv_b, fname, mime, meta = eng.excel_to_csv(html_bytes, sheet_name="Tabel_1", delimiter=";")
        text = csv_b.decode("utf-8-sig")
        self.assertIn("1;Budi", text)

    def test_excel_xml_spreadsheet_disguised(self):
        xml_bytes = (
            b'<?xml version="1.0"?>'
            b'<Workbook xmlns="urn:schemas-microsoft-com:office:spreadsheet"'
            b' xmlns:ss="urn:schemas-microsoft-com:office:spreadsheet">'
            b'<Worksheet ss:Name="SheetLaporan"><Table>'
            b'<Row><Cell><Data ss:Type="String">Header</Data></Cell></Row>'
            b'<Row><Cell><Data ss:Type="String">Data123</Data></Cell></Row>'
            b'</Table></Worksheet></Workbook>'
        )
        info = eng.inspect_excel(xml_bytes)
        self.assertEqual(info["engine"], "xml")
        self.assertIn("SheetLaporan", info["sheet_names"])

        csv_b, fname, mime, meta = eng.excel_to_csv(xml_bytes, sheet_name="SheetLaporan", delimiter=",")
        text = csv_b.decode("utf-8-sig")
        self.assertIn("Data123", text)

    def test_images_to_pdf(self):
        # Buat 2 gambar dummy: 1 RGB dan 1 RGBA
        img1 = Image.new("RGB", (200, 150), color="blue")
        b1 = io.BytesIO()
        img1.save(b1, format="JPEG")

        img2 = Image.new("RGBA", (150, 200), color="red")
        b2 = io.BytesIO()
        img2.save(b2, format="PNG")

        images = [("foto1.jpg", b1.getvalue()), ("grafik2.png", b2.getvalue())]
        pdf_bytes, fname, mime, meta = eng.images_to_pdf(
            images, page_size="A4", orientation="auto", margin_mm=10, output_name="gabungan.pdf"
        )
        self.assertEqual(mime, "application/pdf")
        self.assertEqual(fname, "gabungan.pdf")
        self.assertEqual(meta["total_pages"], 2)

        reader = pypdf.PdfReader(io.BytesIO(pdf_bytes))
        self.assertEqual(len(reader.pages), 2)

    def test_pdf_to_images(self):
        # 1. Multi halaman -> Output ZIP
        zip_bytes, fname, mime, meta = eng.pdf_to_images(
            self.pdf_bytes, dpi=72, img_format="png", base_filename="dokumen_test"
        )
        self.assertEqual(mime, "application/zip")
        self.assertTrue(fname.endswith(".zip"))
        self.assertEqual(meta["total_pages"], 3)
        self.assertTrue(meta["is_zip"])

        with zipfile.ZipFile(io.BytesIO(zip_bytes)) as zf:
            names = zf.namelist()
            self.assertEqual(len(names), 3)
            self.assertTrue(any(n.endswith(".png") for n in names))

        # 2. Satu halaman -> Output single image direct download
        w_single = pypdf.PdfWriter()
        w_single.add_blank_page(width=100, height=100)
        buf_single = io.BytesIO()
        w_single.write(buf_single)

        img_b, fname_s, mime_s, meta_s = eng.pdf_to_images(
            buf_single.getvalue(), dpi=72, img_format="jpeg", base_filename="satu_hal"
        )
        self.assertEqual(mime_s, "image/jpeg")
        self.assertFalse(meta_s["is_zip"])
        self.assertTrue(fname_s.endswith(".jpg"))

    def test_compress_pdf(self):
        # Buat PDF yang berisi gambar
        img = Image.new("RGB", (300, 300), color="green")
        b_img = io.BytesIO()
        img.save(b_img, format="JPEG")

        pdf_bytes, _, _, _ = eng.images_to_pdf([("gambar.jpg", b_img.getvalue())])
        orig_len = len(pdf_bytes)

        comp_bytes, fname, mime, meta = eng.compress_pdf(
            pdf_bytes, level="high", base_filename="surat_resmi"
        )
        self.assertEqual(mime, "application/pdf")
        self.assertTrue(fname.endswith(".pdf"))
        self.assertEqual(meta["total_pages"], 1)
        self.assertIn("saved_bytes", meta)
        self.assertIn("saved_percent", meta)

        reader = pypdf.PdfReader(io.BytesIO(comp_bytes))
        self.assertEqual(len(reader.pages), 1)

    def test_rotate_pdf(self):
        # 1. Putar semua halaman 90 derajat
        rot_bytes, fname, mime, meta = eng.rotate_pdf(
            self.pdf_bytes, rotation=90, pages_mode="all", base_filename="surat"
        )
        self.assertEqual(mime, "application/pdf")
        self.assertTrue(fname.endswith(".pdf"))
        self.assertEqual(meta["rotation_applied"], 90)

        reader = pypdf.PdfReader(io.BytesIO(rot_bytes))
        self.assertEqual(len(reader.pages), 3)
        for page in reader.pages:
            self.assertEqual(page.rotation, 90)

        # 2. Putar hanya halaman 2 sebesar 180 derajat
        rot_bytes2, _, _, _ = eng.rotate_pdf(
            self.pdf_bytes, rotation=180, pages_mode="range", range_str="2"
        )
        reader2 = pypdf.PdfReader(io.BytesIO(rot_bytes2))
        self.assertEqual(reader2.pages[0].rotation, 0)
        self.assertEqual(reader2.pages[1].rotation, 180)
        self.assertEqual(reader2.pages[2].rotation, 0)

    def test_watermark_pdf(self):
        wm_bytes, fname, mime, meta = eng.watermark_pdf(
            self.pdf_bytes,
            text="RAHASIA",
            opacity=0.3,
            font_size=40,
            color="red",
            rotation_deg=45,
            pages_mode="all",
            base_filename="berkas_pajak",
        )
        self.assertEqual(mime, "application/pdf")
        self.assertTrue(fname.endswith(".pdf"))
        self.assertEqual(meta["total_pages"], 3)
        self.assertEqual(meta["text"], "RAHASIA")

        reader = pypdf.PdfReader(io.BytesIO(wm_bytes))
        self.assertEqual(len(reader.pages), 3)

    def test_delete_pages_pdf(self):
        # Hapus halaman 2 dari 3 halaman -> sisa 2 halaman
        del_bytes, fname, mime, meta = eng.delete_pages_pdf(
            self.pdf_bytes, pages_to_delete="2", base_filename="laporan"
        )
        self.assertEqual(mime, "application/pdf")
        self.assertTrue(fname.endswith(".pdf"))
        self.assertEqual(meta["remaining_pages"], 2)

        reader = pypdf.PdfReader(io.BytesIO(del_bytes))
        self.assertEqual(len(reader.pages), 2)

        # Coba hapus semua halaman -> harus melempar ValueError
        with self.assertRaises(ValueError):
            eng.delete_pages_pdf(self.pdf_bytes, pages_to_delete="1-3")

    def test_generate_pdf_thumbnails(self):
        res = eng.generate_pdf_thumbnails(self.pdf_bytes, max_pages=2, dpi=50)
        self.assertEqual(res["total_pages"], 3)
        self.assertEqual(res["rendered_pages"], 2)
        self.assertTrue(res["has_more"])
        self.assertEqual(len(res["thumbnails"]), 2)
        self.assertEqual(res["thumbnails"][0]["page"], 1)
        self.assertTrue(res["thumbnails"][0]["thumb"].startswith("data:image/jpeg;base64,"))
        self.assertGreater(res["thumbnails"][0]["width"], 0)
        self.assertGreater(res["thumbnails"][0]["height"], 0)

    def test_rotate_pdf_rotations_map(self):
        # Rotasi halaman 1 = 90 deg, halaman 3 = 270 deg, halaman 2 tidak diputar
        rotations = {1: 90, 3: 270}
        rot_bytes, fname, mime, meta = eng.rotate_pdf(
            self.pdf_bytes, rotations_map=rotations, base_filename="custom_rot"
        )
        self.assertEqual(mime, "application/pdf")
        self.assertTrue(fname.endswith(".pdf"))
        self.assertEqual(meta["rotation_applied"], "custom")

        reader = pypdf.PdfReader(io.BytesIO(rot_bytes))
        self.assertEqual(len(reader.pages), 3)
        self.assertEqual(reader.pages[0].rotation, 90)
        self.assertEqual(reader.pages[1].rotation, 0)
        self.assertEqual(reader.pages[2].rotation, 270)


if __name__ == "__main__":
    unittest.main()


